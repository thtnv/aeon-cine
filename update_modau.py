import docx
import os
import xml.etree.ElementTree as ET

doc_path = r'D:\Downloads\DOANTOTNGHIEP\DATN_CNS_NguyenVanVien_Final.docx'
doc = docx.Document(doc_path)

content_p120 = (
    "Trong kỷ nguyên số hóa và sự bùng nổ của thương mại điện tử, nhu cầu thưởng thức điện ảnh tại các cụm rạp chiếu phim hiện đại ngày càng trở nên phổ biến. "
    "Tại các trung tâm thương mại quy mô lớn như chuỗi siêu thị AEON MALL, lượng khách hàng đến trải nghiệm dịch vụ chiếu phim vào các dịp cuối tuần, ngày lễ và các khung giờ cao điểm là vô cùng đông đúc. "
    "Tuy nhiên, quy trình mua vé và thanh toán truyền thống trực tiếp tại quầy vé thường xuyên dẫn đến tình trạng ùn tắc, người xem phải xếp hàng chờ đợi lâu, gây áp lực lớn lên nhân viên và làm suy giảm đáng kể mức độ hài lòng của khách hàng.\n\n"
    "Khảo sát các nền tảng bán vé phim trực tuyến thành công hàng đầu tại Việt Nam, đặc biệt là hệ sinh thái Galaxy Cinema (https://www.galaxycine.vn/), cho thấy việc số hóa quy trình đặt vé nhanh 4 bước (Chọn phim – Chọn rạp – Chọn ngày – Chọn suất chiếu), hiển thị trực quan sơ đồ phòng chiếu theo thời gian thực và quản lý tài khoản thành viên tích điểm đã giải quyết triệt để các hạn chế trên, mang lại sự tiện ích tối đa cho người dùng và nâng cao hiệu quả vận hành rạp. "
    "Bên cạnh đó, với sự phát triển vượt bậc của Trí tuệ nhân tạo (AI) và các mô hình ngôn ngữ lớn (LLM), việc tích hợp Trợ lý ảo tư vấn thông minh 24/7 và giải pháp kiểm soát vé tự động bằng mã QR Code là bước tiến tất yếu nhằm hiện đại hóa dịch vụ điện ảnh.\n\n"
    "Xuất phát từ yêu cầu thực tiễn đó, đề tài “Xây dựng hệ thống Website đặt vé xem phim trực tuyến tích hợp công nghệ AI tại rạp phim AEON MALL” (thương hiệu rạp AEON CINE) được nghiên cứu và triển khai nhằm cung cấp một giải pháp phần mềm toàn diện, hiện đại, kế thừa quy trình nghiệp vụ tối ưu của Galaxy Cinema, đồng thời tiên phong tích hợp trí tuệ nhân tạo để nâng cao trải nghiệm đặt vé và tự động hóa quy trình quản lý rạp chiếu."
)

content_p122 = (
    "a) Mục tiêu tổng quát:\n"
    "Xây dựng hoàn chỉnh hệ thống Website Đặt Vé Phim Trực Tuyến AEON CINE tích hợp Trợ lý AI và công nghệ quét mã QR Code, số hóa toàn diện quy trình mua vé xem phim từ khâu tra cứu thông tin, chọn suất chiếu, khóa giữ ghế thời gian thực (SeatHold 10 phút), đặt combo bắp nước, áp dụng voucher khuyến mãi, tích lũy điểm thưởng thành viên, thanh toán trực tuyến cho đến khâu kiểm soát vé vào rạp; đồng thời cung cấp công cụ quản trị tập trung, chính xác cho ban quản lý cụm rạp.\n\n"
    "b) Mục tiêu cụ thể:\n"
    "- Đối với Khách hàng (Customer):\n"
    "  + Cung cấp giao diện trực quan, hiện đại chuẩn responsive (tương thích máy tính và thiết bị di động), lấy cảm hứng thiết kế từ nền tảng Galaxy Cinema.\n"
    "  + Cung cấp công cụ Mua vé nhanh (Quick Booking Widget) 4 bước: Chọn phim -> Chọn rạp -> Chọn ngày -> Chọn suất chiếu.\n"
    "  + Hiển thị danh mục phim đang chiếu và sắp chiếu được phân loại độ tuổi rõ ràng (P, K, T13, T16, T18), hỗ trợ xem thông tin chi tiết, trailer phim chất lượng cao.\n"
    "  + Xây dựng sơ đồ phòng chiếu chọn ghế trực quan theo từng loại ghế (Standard, VIP, Sweetbox) với cơ chế khóa giữ ghế thời gian thực (SeatHold trong 10 phút) để tránh xung đột trùng ghế giữa nhiều người dùng.\n"
    "  + Cho phép chọn các gói combo bắp nước đa dạng, áp dụng mã voucher giảm giá và tích lũy điểm thưởng thành viên AEON Member (hạng thẻ Star, G-Star, X-Star).\n"
    "  + Tích hợp Trợ lý ảo AI (sử dụng Google Gemini AI API) hoạt động 24/7 để tư vấn lịch chiếu, giá vé, thể loại và gợi ý phim theo sở thích tự nhiên của người dùng.\n"
    "  + Tự động xuất vé điện tử kèm mã QR Code bảo mật gửi về tài khoản và email khách hàng sau khi thanh toán thành công.\n"
    "- Đối với Nhân viên rạp (Staff):\n"
    "  + Cung cấp giao diện Mobile Web Scanner cho phép nhân viên rạp sử dụng camera thiết bị di động để quét mã QR trên vé điện tử của khách hàng.\n"
    "  + Xác thực tính hợp lệ của vé chỉ trong 1 – 2 giây, tự động cập nhật trạng thái vé sang 'Đã sử dụng' (Checked-in) nhằm ngăn ngừa tình trạng gian lận hoặc tái sử dụng vé.\n"
    "- Đối với Ban quản lý rạp (Admin):\n"
    "  + Xây dựng bảng điều khiển quản trị tập trung (Admin Dashboard) với các biểu đồ thống kê trực quan về doanh thu theo ngày/tháng/năm, số lượng vé bán ra và tỷ lệ lấp đầy ghế.\n"
    "  + Quản lý toàn diện danh mục phim (thông tin, poster, trailer, đạo diễn, diễn viên, thời lượng, phân loại độ tuổi).\n"
    "  + Quản lý hệ thống cụm rạp, phòng chiếu và ma trận sơ đồ ghế ngồi (ghế Thường, VIP, Đôi).\n"
    "  + Quản lý linh hoạt lịch chiếu (suất chiếu theo ngày, giờ chiếu, định dạng phòng chiếu).\n"
    "  + Quản lý menu combo bắp nước, chương trình khuyến mãi (voucher giảm giá phần trăm hoặc số tiền cố định, số lượng phát hành, hạn sử dụng), và quản lý tài khoản người dùng/phân quyền nhân viên."
)

content_p125 = (
    "- Quy trình nghiệp vụ bán vé, xếp lịch chiếu, quản lý phòng chiếu, cơ chế giữ ghế thời gian thực (SeatHold 10 phút) và soát vé bằng mã QR tại các hệ thống rạp chiếu phim hiện đại (tham chiếu mô hình chuẩn của Galaxy Cinema - https://www.galaxycine.vn/).\n"
    "- Các công nghệ và framework lập trình Web Fullstack tiên tiến:\n"
    "  + Giao diện người dùng (Frontend): Thư viện ReactJS, Vite, Tailwind CSS, Lucide Icons, thư viện Axios xử lý gọi API.\n"
    "  + Phía máy chủ (Backend): Nền tảng Node.js, Express.js framework, kiến trúc RESTful API, cơ chế bảo mật xác thực JSON Web Token (JWT).\n"
    "  + Cơ sở dữ liệu: Hệ quản trị cơ sở dữ liệu quan hệ PostgreSQL và công cụ ánh xạ đối tượng Prisma ORM.\n"
    "- Ứng dụng Trí tuệ nhân tạo: Mô hình ngôn ngữ lớn (LLM) thông qua Google Gemini AI API nhằm xây dựng tính năng tư vấn, hỗ trợ khách hàng tự động 24/7.\n"
    "- Kỹ thuật mã hóa, sinh mã và giải mã QR Code phục vụ nghiệp vụ phát hành và soát vé điện tử an toàn."
)

content_p127 = (
    "- Phạm vi không gian và địa bàn áp dụng: Nghiên cứu và triển khai ứng dụng cho hệ thống cụm rạp chiếu phim AEON CINE tọa lạc tại các trung tâm thương mại AEON MALL trên toàn quốc (như Aeon Mall Tân Phú, Bình Tân, Hà Đông, Hải Phòng, Huế...).\n"
    "- Phạm vi nghiệp vụ và chức năng: Đề tài tập trung giải quyết trọn vẹn chu trình đặt vé xem phim trực tuyến (tìm kiếm phim, tra cứu lịch chiếu, chọn ghế realtime, giữ ghế 10 phút, đặt bắp nước combo, áp dụng mã voucher, tích điểm thành viên Star Points, thanh toán hóa đơn, phát hành vé điện tử kèm mã QR); tính năng quét mã QR soát vé tại cửa phòng chiếu dành cho nhân viên rạp; và toàn bộ các mô-đun quản trị rạp chiếu, thống kê doanh thu dành cho Admin.\n"
    "- Phạm vi công nghệ: Xây dựng hệ thống ứng dụng trên nền tảng Web Application, hỗ trợ hiển thị đáp ứng (Responsive) tốt trên cả trình duyệt máy tính cá nhân và các thiết bị di động thông minh (smartphone, tablet).\n"
    "- Giới hạn của đề tài:\n"
    "  + Đề tài tập trung tối ưu hóa trên nền tảng Web Application, chưa phát triển ứng dụng di động độc lập (Native App cho iOS và Android).\n"
    "  + Về phương thức thanh toán: Hệ thống xây dựng quy trình mô phỏng xác thực giao dịch trực tuyến an toàn theo chuẩn thương mại điện tử, chưa kết nối trực tiếp với cổng thanh toán thực tế của các ngân hàng thương mại (như VNPay, MoMo, ZaloPay) do các ràng buộc về mặt pháp lý và tài khoản doanh nghiệp."
)

content_p129 = (
    "Để thực hiện đề tài một cách khoa học, hiệu quả và đạt tính ứng dụng cao, các phương pháp nghiên cứu sau đã được áp dụng:\n"
    "1. Phương pháp nghiên cứu tài liệu lý thuyết:\n"
    "- Thu thập, đọc và tổng hợp các tài liệu chuyên ngành về công nghệ phần mềm, kiến trúc Single Page Application (SPA), thiết kế hệ thống RESTful API chuẩn mực.\n"
    "- Nghiên cứu nguyên lý thiết kế và tối ưu hóa cơ sở dữ liệu quan hệ với PostgreSQL và Prisma ORM; các giải pháp quản lý trạng thái đồng thời (concurrency control) nhằm xử lý bài toán giữ ghế thời gian thực (SeatHold 10 phút), ngăn chặn xung đột đặt trùng ghế.\n"
    "- Nghiên cứu tài liệu kỹ thuật tích hợp AI của Google (Google Gen AI SDK / Gemini 2.5 Flash API) và các giải pháp sinh/quét mã QR Code trong môi trường Web.\n"
    "2. Phương pháp khảo sát và phân tích thực tiễn:\n"
    "- Khảo sát thực tế quy trình đặt vé và soát vé tại các rạp chiếu phim hiện nay; phân tích sâu mô hình giao diện, luồng trải nghiệm người dùng (UX) và các tính năng cốt lõi trên trang web chính thức của Galaxy Cinema (https://www.galaxycine.vn/) – một trong những thương hiệu rạp chiếu phim hàng đầu Việt Nam.\n"
    "- Tiếp thu và chuẩn hóa các quy tắc nghiệp vụ thực tế: luồng mua vé nhanh 4 bước, phân loại phim theo độ tuổi (P, K, T13, T16, T18), hệ thống bắp nước combo, chính sách điểm thưởng thành viên (Star, G-Star, X-Star) để áp dụng vào hệ thống AEON CINE.\n"
    "3. Phương pháp phân tích và thiết kế hệ thống:\n"
    "- Sử dụng phương pháp phân tích thiết kế hướng đối tượng (OOAD) với ngôn ngữ mô hình hóa UML.\n"
    "- Xây dựng các sơ đồ nghiệp vụ thực tế và phần mềm, biểu đồ ca sử dụng (Use Case Diagram), biểu đồ hoạt động (Activity Diagram), biểu đồ lớp (Class Diagram), biểu đồ tuần tự (Sequence Diagram) và sơ đồ thực thể liên kết (ERD) bằng công cụ StarUML.\n"
    "4. Phương pháp thực nghiệm (Lập trình và kiểm thử):\n"
    "- Cài đặt môi trường phát triển, hiện thực hóa các chức năng phần mềm bằng mã nguồn thực tế (ReactJS, Tailwind CSS, Node.js, Express.js, PostgreSQL, Prisma ORM).\n"
    "- Tích hợp Trợ lý ảo AI Google Gemini và mô-đun quét mã QR soát vé.\n"
    "- Tiến hành kiểm thử chức năng (Functional Testing), kiểm thử luồng nghiệp vụ đặt vé xuyên suốt từ khách hàng đến nhân viên soát vé và quản trị viên, kiểm thử tải và cơ chế khóa giữ ghế thời gian thực nhằm đảm bảo tính ổn định và chính xác cao của hệ thống."
)

content_p131 = (
    "Nội dung báo cáo đồ án tốt nghiệp được tổ chức thành các phần và chương chính như sau:\n"
    "- MỞ ĐẦU: Trình bày bối cảnh và tính cấp thiết của đề tài, mục tiêu nghiên cứu (tổng quát và cụ thể), đối tượng và phạm vi nghiên cứu, các phương pháp nghiên cứu được sử dụng và tổng quan cấu trúc của đồ án.\n"
    "- CHƯƠNG 1: CƠ SỞ LÝ THUYẾT: Giới thiệu tổng quan về hệ thống đặt vé xem phim trực tuyến và phân tích mô hình rạp tham chiếu Galaxy Cinema; trình bày cơ sở lý thuyết về các công nghệ được sử dụng trong đề tài, bao gồm công nghệ Frontend (ReactJS, Vite, Tailwind CSS), Backend (Node.js, Express.js, RESTful API), Cơ sở dữ liệu (PostgreSQL, Prisma ORM), công nghệ Trí tuệ nhân tạo (Google Gemini AI API) và giải pháp soát vé điện tử bằng mã QR Code.\n"
    "- CHƯƠNG 2: PHÂN TÍCH THIẾT KẾ HƯỚNG ĐỐI TƯỢNG: Trình bày kết quả khảo sát yêu cầu, mô tả chi tiết hoạt động nghiệp vụ thực tế và nghiệp vụ phần mềm; xác định các tác nhân (Khách hàng, Nhân viên soát vé, Quản trị viên) và yêu cầu chức năng; xây dựng các sơ đồ Use-case tổng quát và chi tiết, kịch bản hoạt động, phác thảo giao diện Wireframe, thiết kế cơ sở dữ liệu (ERD và mô hình vật lý), sơ đồ lớp (Class Diagram) và các sơ đồ tuần tự (Sequence Diagram) cho các kịch bản nghiệp vụ trọng yếu.\n"
    "- CHƯƠNG 3: XÂY DỰNG CHƯƠNG TRÌNH: Giới thiệu công cụ, môi trường phát triển và cấu trúc mã nguồn dự án; trình bày chi tiết kết quả xây dựng giao diện và hoạt động của các chức năng trong hệ thống (Giao diện Trang chủ, Mua vé nhanh 4 bước, Sơ đồ chọn ghế realtime và cơ chế giữ ghế 10 phút, Đặt combo bắp nước & Áp dụng voucher, Trợ lý AI tư vấn phim, Ứng dụng Mobile Web quét QR vé cho nhân viên, và Hệ thống quản trị Admin đa năng).\n"
    "- KẾT LUẬN & HƯỚNG PHÁT TRIỂN: Tổng kết và đánh giá các kết quả đạt được của hệ thống so với các mục tiêu ban đầu đặt ra; phân tích những hạn chế còn tồn tại và đề xuất các hướng nghiên cứu, nâng cấp và phát triển mở rộng hệ thống trong tương lai.\n"
    "- TÀI LIỆU THAM KHẢO: Liệt kê các tài liệu khoa học, sách, bài báo và các nguồn tài liệu kỹ thuật trực tuyến đã được tham khảo trong quá trình thực hiện đồ án."
)

def set_para_content(p, text, style_name='Normal'):
    p.text = ''
    p.style = doc.styles[style_name]
    pPr = p._p.pPr
    if pPr is not None:
        numPr = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr')
        if numPr is not None:
            pPr.remove(numPr)
        ind = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ind')
        if ind is not None:
            pPr.remove(ind)
            
    lines = text.split('\n')
    for i, line in enumerate(lines):
        r = p.add_run(line)
        if i < len(lines) - 1:
            r.add_break()

set_para_content(doc.paragraphs[120], content_p120, 'Normal')
set_para_content(doc.paragraphs[122], content_p122, 'Normal')
set_para_content(doc.paragraphs[125], content_p125, 'Normal')
set_para_content(doc.paragraphs[127], content_p127, 'Normal')
set_para_content(doc.paragraphs[129], content_p129, 'Normal')
set_para_content(doc.paragraphs[131], content_p131, 'Normal')

doc.save(doc_path)
print("Updated successfully:", doc_path)

# Also save copy
copy_path = r'D:\Downloads\DATN_CNS_NguyenVanVien_Final_CapNhat.docx'
doc.save(copy_path)
print("Saved updated copy:", copy_path)
